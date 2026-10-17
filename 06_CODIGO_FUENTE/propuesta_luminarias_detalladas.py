"""Propuesta REV B: dibujo CAD con detalle interior, según captura del usuario.
No abre el DXF de referencia (supera el límite de transferencia), ni crea bloques
definitivos ni modifica planos. Uso: python propuesta_luminarias_detalladas.py
cuadro.xlsx carpeta_salida
"""
from pathlib import Path
import sys,json,math
import openpyxl,fitz
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from propuesta_luminarias_pdf import Sheet,W,H,FONT,INK,BLUE,GREY,EDGE,BLACK

CYAN='#138ab2';LIGHT='#81bbce'


class Detailed(Sheet):
 def footer(self,note='PROPUESTA GRÁFICA DE PROYECTO · REV B · Acabado inspirado en la captura facilitada por el usuario.'):
  self.line(16,280,404,280,EDGE)
  self.txt(16,287,note,8,col=GREY)
  self.txt(365,287,f'REV B / {self.n:02d}',8,True,BLUE)
  self.c.showPage()
 def frame(self,x,y,s,col=CYAN):
  for inset,lw in [(0,.55),(.025*s,.18),(.055*s,.25)]:
   self.rect(x-s/2+inset,y-s/2+inset,s-2*inset,s-2*inset,None,col,lw*min(1,s/26))
 def cross(self,x,y,r,col=CYAN):
  lw=max(.08,min(.25,r*.065))
  self.circle(x,y,r,None,col,lw)
  self.line(x-r*1.38,y,x+r*1.38,y,col,lw)
  self.line(x,y-r*1.38,x,y+r*1.38,col,lw)
 def label(self,x,y,s,t,size,col):
  self.c.setFont('OR',size)
  width=pdfmetrics.stringWidth(t,'OR',size)/mm
  self.txt(x+s/2-s*.09-width,y+s/2-s*.09,t,size,col=col)
 def panel(self,x,y,s,tag='',col=CYAN,led=True):
  self.frame(x,y,s,col)
  # La doble diagonal de esquina y el centro recuperan el lenguaje de la captura.
  a=x-s/2+s*.055;b=y-s/2+s*.055
  for delta in [.18,.24]:
   self.line(a,b+s*delta,a+s*delta,b,col,.32)
  self.cross(x,y,s*.15,col)
  if led:self.label(x,y,s,'LED',max(4.4,s*.72),col)
  if tag:
   self.txt(x-s*.37,y+s*.39,tag,max(4.0,s*.28),True,col)
 def linear(self,x,y,ww,hh=None,col=CYAN,variant=''):
  hh=hh or ww*.11
  self.rect(x-ww/2,y-hh/2,ww,hh,None,col,.48,min(hh*.25,1.6))
  self.rect(x-ww/2+.8,y-hh/2+.65,ww-1.6,hh-1.3,None,col,.2,min(hh*.2,1.1))
  # Difusor continuo, caras de extremo, dos presillas y nervios longitudinales.
  for yy in [-.21,0,.21]:self.line(x-ww*.425,y+hh*yy,x+ww*.425,y+hh*yy,LIGHT if col==CYAN else col,.15)
  for xx in [x-ww*.465,x+ww*.465]:self.line(xx,y-hh*.38,xx,y+hh*.38,col,.3)
  for xx in [x-ww*.32,x+ww*.32]:
   self.rect(xx-ww*.011,y-hh*.54,ww*.022,hh*1.08,None,col,.23,.15)
  self.circle(x-ww*.48,y,hh*.075,None,col,.18)
  if ww>40:self.txt(x-ww*.09,y+hh*.28,'LED',max(5,hh*1.4),True,col)
  if variant:self.txt(x+ww*.21,y+hh*.28,variant,max(4.5,hh*1.05),True,col)
 def circular(self,kind,x,y,d,col=CYAN):
  r=d/2
  if kind=='wall':
   for rr,lw in [(r,.50),(.94*r,.18),(.87*r,.25)]:self.circle(x,y,rr,None,col,lw)
   for angle in [0,90,180,270]:
    a=math.radians(angle);self.circle(x+math.cos(a)*r*.90,y+math.sin(a)*r*.90,r*.024,None,col,.18)
   self.cross(x,y,r*.33,col)
   self.label(x,y,d*.82,'LED',max(5,d*.66),col)
  elif kind=='recess':
   for rr,lw in [(r,.5),(.87*r,.22),(.68*r,.32),(.62*r,.15)]:self.circle(x,y,rr,None,col,lw)
   for xx in [x-r*1.05,x+r*1.05]:
    self.line(xx,y-r*.17,xx,y+r*.17,col,.25)
   self.cross(x,y,r*.24,col)
   self.txt(x-r*.28,y+r*.51,'LED',max(4.4,d*.47),True,col)
  else:
   for rr,lw in [(r,.55),(.92*r,.35),(.82*r,.22),(.65*r,.25)]:self.circle(x,y,rr,None,col,lw)
   for angle in [30,150,270]:
    a=math.radians(angle);xx=x+math.cos(a)*r*.88;yy=y+math.sin(a)*r*.88
    self.line(xx-r*.04,yy,xx+r*.04,yy,col,.22)
   self.cross(x,y,r*.24,col)
   self.txt(x-r*.28,y+r*.51,'LED',max(4.4,d*.47),True,col)
 def fixture(self,s,x,y,size=27,col=CYAN,tag=True):
  code=s['code']
  if code.startswith('L-01'):self.linear(x,y,size,col=col,variant='15W' if code=='L-01-15' else '1800' if code=='L-01-18' else '')
  elif code.startswith('L-03'):
   t='DALI' if code in ('L-03-D','L-03-DA') else 'IP65' if code=='L-03-IP65' else '90' if code=='L-03-90' else ''
   self.panel(x,y,size,t if tag else '',col)
  else:self.circular('wall' if code=='L-02' else 'recess' if code=='L-04B' else 'surface',x,y,size,col)
 def polygon(self,points,fill,col=INK,lw=.25):
  p=self.c.beginPath();p.moveTo(points[0][0]*mm,(H-points[0][1])*mm)
  for x,y in points[1:]:p.lineTo(x*mm,(H-y)*mm)
  p.close();self.c.setStrokeColor(HexColor(col));self.c.setFillColor(HexColor(fill));self.c.setLineWidth(lw*mm)
  self.c.drawPath(p,stroke=1,fill=1)


def main(excel,out):
 out=Path(out);out.mkdir(parents=True,exist_ok=True)
 pdfmetrics.registerFont(TTFont('OR',FONT+'OpenSans-Regular.ttf'))
 pdfmetrics.registerFont(TTFont('OB',FONT+'OpenSans-Semibold.ttf'))
 w=openpyxl.load_workbook(excel,data_only=True);specs=[]
 for r in range(5,17):
  s=w['Leyenda'];code=s.cell(r,1).value
  specs.append(dict(code=code,power=s.cell(r,3).value,lm=s.cell(r,4).value,ip=s.cell(r,7).value,desc=s.cell(r,2).value))
 by={s['code']:s for s in specs}
 pdf=out/'PROPUESTA_LUMINARIAS_DETALLADAS_REV_B.pdf'
 c=canvas.Canvas(str(pdf),pagesize=(W*mm,H*mm),pageCompression=1)
 c.setTitle('Luminarias MCF Piura · Propuesta detallada REV B')
 sh=Detailed(c)
 sh.header('Una propuesta con más trabajo de dibujo','12 tipos del cuadro · Marco, difusor, centro y detalles interiores · Referencia de acabado: captura del usuario')
 for i,s in enumerate(specs):
  x=16+i%4*98;y=55+i//4*71
  sh.rect(x,y,94,66,None,EDGE,.23,1.5)
  sh.txt(x+5,y+9,s['code'],14,True)
  sh.para(x+5,y+13,84,s['desc'],7.3,col=GREY)
  size=72 if s['code']=='L-01' else 47 if s['code'].startswith('L-01-') else 26
  sh.fixture(s,x+47,y+36,size)
  sh.txt(x+5,y+59,f"{s['power']} W · {s['lm']} lm · {s['ip']}",8,col=INK)
 sh.para(19,270,380,'Las marcas interiores expresan el estilo gráfico del bloque. Los códigos y atributos conservan la identificación de cada tipo. El dibujo no representa una fabricación verificada ni sustituye la simbología eléctrica DGE.',8)
 sh.footer()

 # Detalle de panel: lectura directa del ejemplo recibido.
 sh.header('Panel LED / interpretación de tu referencia','Marco trabajado, esquina diagonal, círculo con ejes e identificación LED. El código del tipo queda fuera del dibujo.')
 sh.rect(17,56,183,137,None,EDGE,.25,2)
 sh.txt(26,68,'PROPUESTA PARA L-03',12,True,BLUE)
 sh.fixture(by['L-03'],107,125,86)
 sh.txt(143,77,'L-03',13,True,BLUE)
 sh.txt(216,68,'VOLUMEN ORIENTATIVO DEL EQUIPO',11,True,BLUE)
 pts=[(232,113),(302,83),(385,126),(315,158)]
 sh.polygon([pts[0],pts[3],(315,162),(232,117)],'#b9c8d0')
 sh.polygon([pts[3],pts[2],(385,130),(315,162)],'#8eabb9')
 sh.polygon(pts,'#edf2f5')
 sh.polygon([(239,113),(302,86),(377,126),(315,154)],'#ffffff',GREY,.18)
 sh.para(216,175,181,'Volumen ilustrativo de un panel con marco y difusor. Geometría elaborada para esta propuesta; no es fotografía ni modelo de fabricación.',9,col=GREY)
 descriptions=[('MARCO','Tres líneas con distinto peso para dar profundidad.'),('ESQUINA','Doble diagonal como la referencia que enviaste.'),('CENTRO','Círculo y ejes; centro geométrico del futuro bloque.'),('IDENTIDAD','LED y distintivo de variante con espacio reservado.')]
 for i,(title,note) in enumerate(descriptions):
  x=19+i*98
  sh.txt(x,215,title,11,True,BLUE)
  sh.para(x,222,90,note,10)
 sh.para(19,255,380,'Propuesta: mantener este lenguaje en toda la familia y diferenciar L-03, L-03-D, L-03A, L-03-DA, L-03-IP65 y L-03-90 mediante sus códigos y datos. DALI e IP son identificadores de proyecto.',9)
 sh.footer()

 sh.header('Lineales / carcasa, difusor y presillas','Las versiones corta y larga conservan un lenguaje común. Se amplía el detalle para revisar el acabado.')
 for i,code in enumerate(['L-01','L-01-15','L-01-18']):
  yy=69+i*62;s=by[code]
  sh.txt(20,yy,code,16,True)
  sh.txt(20,yy+10,f"{s['power']} W / {s['lm']} lm",10,col=BLUE)
  sh.linear(205,yy+8,188 if i==0 else 113,15 if i==0 else 12,variant='15W' if i==1 else '1800' if i==2 else '')
  sh.txt(331,yy+3,'Extremos definidos',9)
  sh.txt(331,yy+10,'Difusor continuo',9)
  sh.txt(331,yy+17,'Presillas orientativas',9)
  sh.line(20,yy+34,400,yy+34,EDGE)
 sh.para(20,255,380,'Las presillas y líneas interiores son detalles gráficos de la propuesta. La longitud nominal distingue la lineal larga y la corta; el código evita confundir las dos versiones de 15 W. Los anchos y detalles se ajustarán con el modelo definitivo.',9)
 sh.footer()

 sh.header('Circulares / una identidad por montaje','Aplique, downlight empotrado y downlight adosado: aros, difusor, centro y perfil de montaje.')
 for i,code in enumerate(['L-02','L-04B','L-04C']):
  x=16+i*132;s=by[code]
  sh.rect(x,56,124,211,None,EDGE,.25,1.5)
  sh.txt(x+6,69,code,18,True)
  sh.txt(x+6,79,'Aplique de pared' if i==0 else 'Downlight empotrado' if i==1 else 'Downlight de superficie',10,col=BLUE)
  sh.fixture(s,x+62,123,52)
  sh.txt(x+6,162,'PERFIL ORIENTATIVO',9,True,GREY)
  if i==0:
   sh.line(x+28,176,x+96,176,INK,.65)
   sh.rect(x+43,177,38,8,None,CYAN,.4,2)
   for off in [2,4]:sh.line(x+45,177+off,x+79,177+off,LIGHT,.18)
  elif i==1:
   sh.line(x+25,185,x+99,185,GREY,.25,[2,1])
   sh.rect(x+43,172,38,12,None,CYAN,.35)
   sh.rect(x+38,184,48,2,None,CYAN,.45)
   sh.rect(x+47,185,30,1,None,CYAN,.2)
  else:
   sh.line(x+25,172,x+99,172,INK,.6)
   sh.rect(x+41,174,42,12,None,CYAN,.4,1)
   sh.line(x+43,184,x+81,184,LIGHT,.25)
  note=['Aro perimetral, puntos de detalle y centro. La vista circular es frontal; el perfil explica la posición junto al muro.',
   'Aro de empotramiento y reflector/difusor interior. El perfil deja el cuerpo por encima del falso cielo raso.',
   'Bisel más marcado y cuerpo adosado. El perfil distingue su espesor bajo la losa.'][i]
  sh.para(x+6,201,111,note,10)
  sh.para(x+6,245,111,'Detalles orientativos de dibujo; el código identifica el montaje y el equipo.',8.5,col=GREY)
 sh.footer()

 sh.header('Legibilidad en plano / detalle completo y reducido','Se conserva el código a 2,5 mm impresos. Los gráficos se dimensionan para lectura; no representan la huella física del equipo.')
 for i,den in enumerate([50,75,100]):
  x=16+i*132;ww=5600/den;hh=3700/den;xx=x+(124-ww)/2;yy=83
  sh.txt(x+5,64,f'ESCALA 1:{den}',15,True,BLUE)
  sh.rect(xx,yy,ww,hh,None,INK,.25)
  for px,py,code in [(1400,1100,'L-03'),(4200,1100,'L-03-D'),(1400,2600,'L-03-IP65'),(4200,2600,'L-03-DA')]:
   cx=xx+px/den;cy=yy+py/den
   if den==50:sh.fixture(by[code],cx,cy,11,BLACK)
   else:
    # Suprimir texto interior y segunda diagonal cuando el tamaño lo exige.
    sh.frame(cx,cy,8,BLACK);sh.cross(cx,cy,1.25,BLACK)
    sh.line(cx-3.55,cy-1.9,cx-1.9,cy-3.55,BLACK,.18)
   sh.txt(cx-6,cy+9,code,2.5*mm,col=INK)
  sh.para(x+5,179,114,'Detalle completo' if den==50 else 'Detalle reducido para impresión',10,col=BLUE)
  sh.para(x+5,190,114,'Marco, centro y código mantienen la identidad de la familia. Las líneas finas se simplifican al imprimir pequeño.',9,col=GREY)
 sh.line(19,221,401,221,EDGE)
 sh.txt(20,234,'CRITERIO DE ESTA REVISIÓN',11,True,BLUE)
 sh.para(20,241,380,'El acabado responde a la captura compartida: dibujo vectorial con marco, difusor, centro, ejes y detalles de familia. Son marcas gráficas de proyecto; la simbología DGE permanece como referencia eléctrica separada. Esta revisión entrega ideas en PDF y deja pendiente tu elección del acabado antes de construir los bloques definitivos.',9)
 sh.footer('Referencia visual: captura en el chat. El DXF adjunto no se pudo transferir por superar 32 MB. Datos técnicos: Excel del usuario.')
 c.save()
 doc=fitz.open(pdf);assert len(doc)==5
 for i,page in enumerate(doc):
  assert len(page.get_text())>250
  page.get_pixmap(dpi=110,alpha=False).save(out/f'LUMINARIAS_REV_B_{i+1:02d}.png')
 for s in specs:assert s['code'] in '\n'.join(p.get_text() for p in doc)
 (out/'REFERENCIAS_REV_B.json').write_text(json.dumps(dict(tipos=specs,referencia_estilo='Captura del usuario: marco, esquina doble, circulo con ejes y LED',
   dxf_referencia_leido=False,motivo='Archivo supera limite de transferencia de 32 MiB',planos_modificados=False,bloques_definitivos_creados=False),ensure_ascii=False,indent=2))
 print(json.dumps({'pdf':str(pdf),'paginas':len(doc),'tipos':len(specs),'bytes':pdf.stat().st_size},ensure_ascii=False))


if __name__=='__main__':main(*sys.argv[1:])
