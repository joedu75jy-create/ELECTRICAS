"""Presentación REV F de los doce tipos, sin códigos debajo de los símbolos.
Reutiliza exactamente el dibujo L-03 de la revisión C aprobada para presentación.
No crea bloques definitivos. Uso: python presentar_todas_luminarias.py cuadro.xlsx salida
"""
from pathlib import Path
import sys,json
import openpyxl,fitz
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from propuesta_l03_sin_rotulo import Drawing
from propuesta_luminarias_detalladas import Detailed
from propuesta_luminarias_pdf import W,H,FONT,INK,BLUE,GREY,EDGE

BODY='#10161a'


class Catalog(Detailed):
 def footer(self,note='Códigos únicamente en cabeceras de ficha · Símbolos sin rótulo exterior · Propuesta gráfica de proyecto.'):
  self.line(16,280,404,280,EDGE)
  self.txt(16,287,note,8,col=GREY)
  self.txt(363,287,f'REV F / {self.n:02d}',8,True,BLUE)
  self.c.showPage()
 def linear(self,x,y,ww,hh=None,col=BODY,variant=''):
  hh=hh or ww*.11
  super().linear(x,y,ww,hh,col,variant='',led=False)
  if ww>40:
   led_pt=max(5,hh*1.4)
   led_width=pdfmetrics.stringWidth('LED','OR',led_pt)/mm
   self.txt(x-led_width/2,y+hh*.28,'LED',led_pt,False,col)
  if variant:
   pt=max(4.5,hh*.9)
   width=pdfmetrics.stringWidth(variant,'OB',pt)/mm
   lx=x+ww*.295-width
   self.rect(lx-.45,y-hh*.08,width+.9,hh*.5,fill='#ffffff',col='#ffffff',lw=0)
   self.txt(lx,y+hh*.28,variant,pt,True,col)
 def circular(self,kind,x,y,d,col=BODY):
  super().circular(kind,x,y,d,col,led=False,center_scale=.20)
  # Una unica regla para las tres circulares, tanto en vista general como detalle.
  pt=d*.15*mm
  width=pdfmetrics.stringWidth('LED','OR',pt)/mm
  self.txt(x-width/2,y+d*.26,'LED',pt,False,col)
 def lamp(self,code,x,y,size):
  if code.startswith('L-03'):
   d=Drawing(self.c,H)
   d.panel(x-size/2,y-size/2,size,BODY)
   if code in ('L-03A','L-03-DA'):
    # Esquina adicional: marca de proyecto para distinguir el flujo reducido.
    left=x-size/2;top=y-size/2
    for edge in [.29,.35]:
     d.line(left+size*.915,top+size*edge,left+size*(1-edge),top+size*.085,BODY,max(.075,min(.26,size*.011)))
   tag={'L-03-D':'DALI','L-03-DA':'DALI','L-03-IP65':'IP65','L-03-90':'IRC90'}.get(code)
   if tag:d.text(x-size*.35,y+size*.34,tag,size*.075*mm,BODY)
  elif code.startswith('L-01'):
   self.linear(x,y,size,col=BODY,variant='15W-A' if code=='L-01-15' else '15W-B' if code=='L-01-18' else '')
  else:
   self.circular('wall' if code=='L-02' else 'recess' if code=='L-04B' else 'surface',x,y,size,BODY)


def main(excel,out):
 out=Path(out);out.mkdir(parents=True,exist_ok=True)
 for name,file in [('OR','OpenSans-Regular.ttf'),('R','OpenSans-Regular.ttf'),('OB','OpenSans-Semibold.ttf'),('B','OpenSans-Semibold.ttf')]:
  pdfmetrics.registerFont(TTFont(name,FONT+file))
 w=openpyxl.load_workbook(excel,data_only=True);specs=[]
 for r in range(5,17):
  s=w['Leyenda'];specs.append(dict(code=s.cell(r,1).value,desc=s.cell(r,2).value,power=s.cell(r,3).value,
    lm=s.cell(r,4).value,ip=s.cell(r,7).value,mount=s.cell(r,8).value))
 assert len(specs)==12 and len({s['code'] for s in specs})==12
 by={s['code']:s for s in specs}
 fn=out/'TODAS_LUMINARIAS_SIN_ROTULOS_REV_F.pdf'
 c=canvas.Canvas(str(fn),pagesize=(W*mm,H*mm),pageCompression=1)
 c.setTitle('MCF Piura · Los 12 tipos de luminarias sin códigos exteriores · REV F')
 sh=Catalog(c)
 sh.header('Las 12 luminarias / presentación completa','Acabado del panel L-03 aprobado · Sin códigos debajo de los dibujos · LED conservado en el interior')
 for i,s in enumerate(specs):
  x=16+i%4*98;y=55+i//4*71
  sh.rect(x,y,94,66,None,EDGE,.2,1.5)
  sh.txt(x+5,y+9,s['code'],14,True)
  sh.para(x+5,y+13,84,s['desc'],7.3,col=GREY)
  size=72 if s['code']=='L-01' else 47 if s['code'].startswith('L-01-') else 27
  sh.lamp(s['code'],x+47,y+36,size)
  sh.txt(x+5,y+59,f"{s['power']} W · {s['lm']} lm · {s['ip']}",8,col=INK)
 sh.para(19,270,380,'Los códigos de arriba identifican las fichas de esta leyenda; no acompañan al símbolo en el plano. Las variantes conservan marcas interiores de proyecto para facilitar su reconocimiento.',8)
 sh.footer()

 sh.header('Paneles / seis variantes de la misma familia','Marco, esquina, círculo y ejes del L-03 aprobado. Las variantes añaden una marca interior y mantienen el perímetro libre.')
 notes={'L-03':'Versión estándar: reproduce el dibujo de la revisión C.',
   'L-03-D':'DALI interior: versión con driver regulable.',
   'L-03A':'Esquina adicional: propuesta para reconocer flujo reducido.',
   'L-03-DA':'Esquina adicional y DALI: flujo reducido con regulación.',
   'L-03-IP65':'IP65 interior: identificación de la versión hermética.',
   'L-03-90':'IRC90 interior: identificación de alta reproducción cromática.'}
 for i,s in enumerate([v for v in specs if v['code'].startswith('L-03')]):
  x=16+i%3*132;y=56+i//3*108
  sh.rect(x,y,124,102,None,EDGE,.2,1.5)
  sh.txt(x+6,y+11,s['code'],15,True)
  sh.txt(x+6,y+20,f"{s['power']} W · {s['lm']} lm · {s['ip']}",9,col=BLUE)
  sh.lamp(s['code'],x+37,y+53,44)
  sh.para(x+67,y+36,50,notes[s['code']],9)
  sh.para(x+6,y+86,111,'Módulo nominal 60 × 60 cm · Empotrado en falso cielo raso.',8,col=GREY)
 sh.footer('Esquina adicional, DALI, IP65 e IRC90 son identificadores gráficos de proyecto; no representan figuras normativas nuevas.')

 sh.header('Lineales / variantes de 15 W identificadas como A y B','15W-A: 2.250 lm · 15W-B: 1.800 lm · La marca interior usa potencia y variante en ambos dibujos.')
 for i,code in enumerate(['L-01','L-01-15','L-01-18']):
  yy=69+i*62;s=by[code]
  sh.txt(20,yy,code,16,True)
  sh.txt(20,yy+10,f"{s['power']} W / {s['lm']} lm / IP65",10,col=BLUE)
  sh.lamp(code,205,yy+8,188 if i==0 else 113)
  sh.txt(331,yy+3,{'L-01-15':'Interior: 15W-A','L-01-18':'Interior: 15W-B'}.get(code,'Marco y extremos definidos'),9)
  sh.txt(331,yy+10,'Flujo indicado en la ficha' if code!='L-01' else 'Difusor continuo',9)
  sh.txt(331,yy+17,'Presillas orientativas',9)
  sh.line(20,yy+34,400,yy+34,EDGE)
 sh.para(20,255,380,'Correspondencia de leyenda: 15W-A = L-01-15, 15 W y 2.250 lm; 15W-B = L-01-18, 15 W y 1.800 lm. A/B son identificadores de proyecto. Se conservan los códigos originales del Excel y el acabado del dibujo.',9)
 sh.footer()

 sh.header('Circulares / aplique y downlights','LED uniforme: misma tipografía regular, tamaño relativo y centrado horizontal en las tres luminarias.')
 for i,code in enumerate(['L-02','L-04B','L-04C']):
  x=16+i*132;s=by[code]
  sh.rect(x,56,124,211,None,EDGE,.2,1.5)
  sh.txt(x+6,69,code,18,True)
  sh.txt(x+6,80,f"{s['power']} W / {s['lm']} lm / {s['ip']}",9,col=BLUE)
  sh.lamp(code,x+62,124,53)
  sh.txt(x+6,164,'PERFIL ORIENTATIVO DE MONTAJE',8.5,True,GREY)
  if i==0:
   sh.line(x+25,177,x+99,177,INK,.6)
   sh.rect(x+43,178,38,8,None,BODY,.35,2)
   sh.line(x+46,184,x+78,184,BODY,.18)
  elif i==1:
   sh.line(x+25,185,x+99,185,GREY,.25,[2,1])
   sh.rect(x+43,173,38,11,None,BODY,.3)
   sh.rect(x+38,184,48,2,None,BODY,.4)
  else:
   sh.line(x+25,173,x+99,173,INK,.6)
   sh.rect(x+42,175,40,12,None,BODY,.35,1)
   sh.line(x+44,185,x+80,185,BODY,.18)
  descriptions=['Aplique circular de pared. Se muestra la cara frontal; el perfil ilustra su ubicación junto al muro.',
   'Downlight empotrado: aro exterior y difusor interior; cuerpo alojado sobre el falso cielo raso.',
   'Downlight de superficie: bisel marcado y cuerpo situado bajo la losa.']
  sh.para(x+6,201,111,descriptions[i],10)
  sh.para(x+6,244,111,'La ficha identifica el tipo; el símbolo que se repetirá en planta conserva solo su dibujo.',8.5,col=GREY)
 sh.footer('Datos del Excel recibido · Dibujos vectoriales de proyecto, orientativos · No se modifican planos ni bloques definitivos.')
 c.save()
 d=fitz.open(fn);assert len(d)==4
 text='\n'.join(p.get_text() for p in d)
 for s in specs:assert s['code'] in text
 for i,page in enumerate(d):
  assert abs(page.rect.width-W*mm)<.01
  page.get_pixmap(dpi=110,alpha=False).save(out/f'TODAS_LUMINARIAS_REV_F_{i+1:02d}.png')
 # El gráfico L-03 básico reutiliza Drawing.panel de la revisión aprobada.
 (out/'REV_F_DATOS.json').write_text(json.dumps(dict(tipos=specs,led_circulares={'fuente':'OpenSans-Regular','tamano_mm_por_diametro':.15,'posicion_baseline_por_diametro':.26,'alineacion':'centrado horizontal'},identificacion_interior={'L-01-15':'15W-A','L-01-18':'15W-B'},base_panel='Drawing.panel de propuesta_l03_sin_rotulo.py, REV C',
   codigos_en_simbolos=False,codigos_en_cabeceras_de_ficha=True,paginas=4,planos_modificados=False,bloques_definitivos_creados=False),ensure_ascii=False,indent=2))
 print(json.dumps({'pdf':str(fn),'paginas':len(d),'tipos':len(specs),'bytes':fn.stat().st_size}))


if __name__=='__main__':main(*sys.argv[1:])
